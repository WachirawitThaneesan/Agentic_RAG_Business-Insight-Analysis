"""Final academic artifacts from one preserved manifest. Purely local authoring."""
from pathlib import Path
import json, hashlib, base64, sys, re
from html import escape
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'artifacts'
OUT.mkdir(exist_ok=True)
M = json.loads((ROOT / 'RESULT_MANIFEST_v2.json').read_text(encoding='utf-8'))
ROWS = json.loads((ROOT / 'RESULTS_BY_QUESTION.json').read_text(encoding='utf-8'))
BANK = json.loads((ROOT / 'SOURCE_REPORTS.json').read_text(encoding='utf-8'))
COST = json.loads((ROOT / 'FILTERED_COST_COMPONENTS_v1.json').read_text(encoding='utf-8'))
QA = json.loads((ROOT / 'SOURCE_SPOT_CHECK_INDEX.json').read_text(encoding='utf-8'))
S0 = M['systems']['simple_dense_rag']
S1 = M['systems']['full_app']
TITLE = 'AI-Driven OCR and Web Scraping System for Business Insight Analysis'
TH_TITLE = 'ระบบ AI สำหรับ OCR และ Web Scraping เพื่อวิเคราะห์ข้อมูลเชิงลึกทางธุรกิจ'
MEMBERS = [('นายจักรภัทร พนมวัฒนศักดิ์', '66070026', '66070026@kmitl.ac.th'), ('นายวชิรวิทย์ ธนีสันต์', '66070176', '66070176@kmitl.ac.th')]
ADVISOR = 'ดร.ณัฏฐ์ ดิลกธนากุล'
FACULTY = 'คณะเทคโนโลยีสารสนเทศ'
PROGRAMME = 'เทคโนโลยีปัญญาประดิษฐ์'
INSTITUTION = 'สถาบันเทคโนโลยีพระจอมเกล้าเจ้าคุณทหารลาดกระบัง'
TERM = 'ภาคการศึกษาที่ 2 ปีการศึกษา 2568'
LOGO = ROOT / 'institution_logo_reference.png'

def pct(x):
    return f'{100 * x:.2f}%'

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def svg_text(x, y, text, size=26, fill='#203541', weight=400, anchor='start', thai=False):
    return f'<text x="{x}" y="{y}" font-family="{('Tahoma' if thai else 'Arial')}" font-size="{size}" fill="{fill}" font-weight="{weight}" text-anchor="{anchor}">{escape(str(text))}</text>'

def svg_lines(x, y, lines, size=26, gap=38, fill='#203541', weight=400, thai=False):
    return ''.join((svg_text(x, y + i * gap, t, size, fill, weight, thai=thai) for i, t in enumerate(lines)))

def svg_start(w, h):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}">'

def assets():
    s = svg_start(1200, 300) + '<rect width="1200" height="300" fill="white"/>'
    labels = [('Official reports', ['PDF identity', 'Source URL and SHA-256']), ('Extract and index', ['Typhoon prose / Gemini tables', 'Local bge-m3 embeddings']), ('Retrieve evidence', ['S0 dense search', 'S1 routing and SQL']), ('Answer and inspect', ['Shared Gemini final generator', 'Captured tuples and citations'])]
    for i, (title, lines) in enumerate(labels):
        x = 15 + 300 * i
        s += f'<rect x="{x}" y="55" width="270" height="165" rx="8" fill="#f4f6f6" stroke="#b6c6c9"/>'
        s += svg_text(x + 16, 92, title, 22, weight=700) + svg_lines(x + 16, 132, lines, 16, 28)
        if i < 3:
            s += svg_text(x + 282, 148, '→', 25)
    s += svg_text(600, 267, 'Original source labels and native PDF text are reserved for evaluation, not generation.', 18, anchor='middle') + '</svg>'
    (OUT / 'workflow.svg').write_text(s, encoding='utf-8')
    s = svg_start(1200, 430) + '<rect width="1200" height="430" fill="white"/>'
    s += svg_text(20, 35, 'Recorded primary AI-assisted results', 34, weight=700)
    s += svg_text(20, 68, 'Completed AI-assisted assessments; distinct question, numerical and control denominators.', 20)
    bars = [('Strict complete answers', S0['complete_answer_rate_all498'], S1['complete_answer_rate_all498'], '498 questions per system'), ('Numerical relations', S0['numeric_rate'], S1['numeric_rate'], '356 relations per system'), ('Control abstention', 24 / 32, 28 / 32, '32 controls per system')]
    for i, (label, a, b, den) in enumerate(bars):
        y = 110 + i * 100
        s += svg_text(20, y, label, 26, weight=700) + svg_text(20, y + 30, den, 22)
        for j, (v, c) in enumerate([(a, '#607e8a'), (b, '#176b60')]):
            yy = y - 17 + j * 36
            s += f'<rect x="370" y="{yy}" width="{640 * v}" height="24" fill="{c}"/>' + svg_text(380 + 640 * v, yy + 21, pct(v), 24)
    s += svg_text(370, 406, 'S0 Dense RAG', 26, '#607e8a') + svg_text(650, 406, 'S1 Full App', 26, '#176b60') + '</svg>'
    (OUT / 'primary_results.svg').write_text(s, encoding='utf-8')
    s = svg_start(1684, 2384) + '<rect width="1684" height="2384" fill="#fcfcf9"/>'
    logo = base64.b64encode(LOGO.read_bytes()).decode()
    s += f'<image href="data:image/png;base64,{logo}" x="1460" y="45" width="145" height="145"/>'
    s += svg_text(80, 85, 'THAI BUSINESS REPORTS  /  EXAMINATION PROJECT', 26, '#176b60', 700)
    s += svg_lines(80, 195, ['AI-Driven OCR and Web Scraping', 'System for Business Insight Analysis'], 62, 78, weight=700)
    s += svg_text(80, 329, TH_TITLE, 28, thai=True)
    for i, (name, sid, email) in enumerate(MEMBERS):
        x = 80 + i * 790
        s += svg_text(x, 385, name, 31, weight=700, thai=True) + svg_text(x, 425, f'{sid}  |  {email}', 25)
    s += svg_text(80, 473, 'อาจารย์ที่ปรึกษา  ' + ADVISOR, 27, thai=True)
    s += '<line x1="80" y1="514" x2="1604" y2="514" stroke="#becaca" stroke-width="2"/>'
    s += svg_text(80, 570, 'The question', 38, weight=700) + svg_lines(80, 620, ['Can structured table access and tool routing improve source-bound answers', 'over one-pass dense retrieval on the same Thai annual reports?'], 31, 43)
    s += svg_text(80, 740, 'One corpus. Two measured systems.', 38, weight=700)
    workflow = [('Collect', ['Official PDFs', 'URLs and hashes']), ('Extract', ['OCR and tables', 'Local stored corpus']), ('Retrieve', ['S0 dense search', 'S1 routing + SQL']), ('Inspect', ['Answer tuples', 'Original PDF citations'])]
    for i, (head, lines) in enumerate(workflow):
        x = 80 + i * 390
        s += f'<rect x="{x}" y="783" width="354" height="155" rx="6" fill="#edf2ef"/>'
        s += svg_text(x + 22, 824, head, 31, weight=700) + svg_lines(x + 22, 872, lines, 26, 36)
        if i < 3:
            s += svg_text(x + 365, 865, '→', 28, '#176b60')
    s += svg_text(80, 997, 'Final frozen run', 38, weight=700)
    s += svg_text(80, 1047, '4 reports   ·   1,431 / 1,447 pages indexed   ·   6,928 searchable chunks', 30)
    s += svg_text(80, 1090, '498 answerable questions + 32 controls per system  |  1,060 first outcomes', 29)
    s += svg_text(740, 1155, 'S0 Dense RAG', 28, '#607e8a', 700) + svg_text(1170, 1155, 'S1 Full App', 28, '#176b60', 700)
    resultrows = [('Strict complete answers', f'{pct(S0['complete_answer_rate_all498'])}  (79/498)', f'{pct(S1['complete_answer_rate_all498'])}  (60/498)'), ('Numerical relations', f'{pct(S0['numeric_rate'])}  (119/356)', f'{pct(S1['numeric_rate'])}  (106/356)'), ('Control abstention', '75.00%  (24/32)', '87.50%  (28/32)'), ('Response time p50 / p95', '7.19 / 19.90 s', '13.01 / 35.55 s')]
    for i, (a, b, c) in enumerate(resultrows):
        y = 1223 + i * 74
        s += svg_text(80, y, a, 30, weight=700) + svg_text(740, y, b, 30) + svg_text(1170, y, c, 30)
        s += f'<line x1="80" y1="{y + 25}" x2="1604" y2="{y + 25}" stroke="#d5dddd"/>'
    s += svg_text(80, 1571, 'Measured conclusion', 38, weight=700)
    s += svg_lines(80, 1623, ['Full App did not outperform Dense RAG on strict completion or numerical relations.', 'It abstained more appropriately on controls and took longer to answer.'], 31, 44)
    s += svg_text(80, 1743, 'Why source context matters', 35, weight=700)
    s += svg_lines(80, 1792, ['TISCO Bank cash: 818,057 thousand baht on PDF page 128; rounded 818 million', 'on page 266. Entity, year, scale and requested precision must stay attached.'], 28, 40)
    s += svg_text(80, 1898, 'Evidence limits', 35, weight=700)
    s += svg_lines(80, 1945, ['1,060 / 1,060 AI-assisted assessments complete; no independent human confirmation.', '34 consistency adjudications; 3 defective source labels disclosed with paired sensitivity.', 'Quality goal >80% remains unmet. Strict scores are not certified benchmark accuracy.'], 27, 40)
    s += svg_text(80, 2095, 'Cost transparency', 35, weight=700)
    s += svg_lines(80, 2141, ['Known uncached Gemini token components: ingestion $3.44; S0 queries $4.07; S1 $5.61.', 'These are partial estimates, not invoices. Missing usage, thinking, cache and other costs remain unknown.'], 26, 37)
    s += '<line x1="80" y1="2207" x2="1604" y2="2207" stroke="#becaca" stroke-width="2"/>'
    s += svg_text(80, 2258, FACULTY + '  |  สาขา' + PROGRAMME, 27, thai=True)
    s += svg_text(80, 2302, INSTITUTION + '  |  ' + TERM, 25, thai=True)
    s += svg_text(80, 2345, 'Sources: KCE, OSP, HMPRO and TISCO Bank One Reports 2025; frozen result manifest and source QA.', 22, '#576974') + '</svg>'
    (OUT / 'project_poster_A1.svg').write_text(s, encoding='utf-8')
    (OUT / 'project_poster_A1.html').write_text('<!doctype html><html><meta charset="utf-8"><title>' + TITLE + '</title><style>@page{size:594mm 841mm;margin:0}html,body{margin:0;padding:0;width:594mm;height:841mm}svg{display:block;width:594mm;height:841mm}</style>' + s + '</html>', encoding='utf-8')
    print('Controlled vector artwork and editable A1 poster saved.')

def report():
    d = Document()
    sec = d.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.85)
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.different_first_page_header_footer = True
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        st = d.styles[name]
        st.font.name = 'Times New Roman'
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.size = Pt(11.5 if name == 'Normal' else 18 if name == 'Title' else 15 if name == 'Heading 1' else 12 if name.startswith('Heading') else 10)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.line_spacing = 1.12
        fonts = st.element.get_or_add_rPr().rFonts
        for theme in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + theme), None)
        for border in list(st.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    d.styles['Normal'].paragraph_format.first_line_indent = Inches(0.2)
    for name in ['Front Heading', 'Table Caption', 'Figure Caption']:
        st = d.styles.add_style(name, 1)
        st.base_style = d.styles['Heading 1' if name == 'Front Heading' else 'Caption']
        st.font.color.rgb = RGBColor(0, 0, 0)
        if name == 'Front Heading':
            el = OxmlElement('w:outlineLvl')
            el.set(qn('w:val'), '9')
            st.element.get_or_add_pPr().append(el)
    for name in ['TOC 1', 'TOC 2']:
        st = d.styles[name] if name in d.styles else d.styles.add_style(name, 1)
        st.base_style = d.styles['Normal']
        st.font.name = 'Times New Roman'
        st.font.size = Pt(11)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(1)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.first_line_indent = Inches(0)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    d.core_properties.title = TITLE
    d.core_properties.author = '66070026 and 66070176'
    d.core_properties.subject = 'Final examination report with preserved AI-assisted evaluation and limitations'

    def p(t='', center=False, thai=False, size=None):
        if not thai and len(t) > 140 and ('http' not in t) and ('SHA-256:' not in t):
            protected = {}

            def protect(match):
                key = 'PROTECTEDTOKEN' + chr(65 + len(protected))
                protected[key] = match.group()
                return key
            t = re.sub('\\b(?:[a-f0-9]{12,64}|(?:KCE|HMPRO|OSP|TISCO)\\d{2}q\\d{2}|S[01])\\b', protect, t)
            t = re.sub('([A-Za-z])(?=\\d)', '\\1 ', t)
            t = re.sub('(\\d)(?=[A-Za-z])', '\\1 ', t)
            t = re.sub('(\\d)(and|or)(?=\\d)', '\\1 \\2 ', t)
            for key, value in protected.items():
                t = t.replace(key, value)
            t = t.replace('bge-m 3', 'bge-m3')
            t = t.replace('A 1', 'A1').replace('DecemberBE', 'December BE')
            t = re.sub('([A-Za-z]),(?=\\d)', '\\1, ', t)
            t = re.sub(';(?=\\S)', '; ', t)
        follows_table = d._body._body[-2].tag == qn('w:tbl') if len(d._body._body) > 1 else False
        a = d.add_paragraph()
        a.paragraph_format.first_line_indent = Inches(0) if center else Inches(0.2)
        if follows_table:
            a.paragraph_format.space_before = Pt(5)
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        r = a.add_run(t)
        if thai:
            r.font.name = 'Tahoma'
            rf = r._element.get_or_add_rPr().rFonts
            rf.set(qn('w:eastAsia'), 'Tahoma')
            rf.set(qn('w:cs'), 'Tahoma')
        if size:
            r.font.size = Pt(size)
        return a

    def page(t, level=1):
        a = d.add_heading(t, level=level)
        a.paragraph_format.page_break_before = True

    def h(t, txt):
        d.add_heading(t, level=2)
        p(txt)

    def cap(t, kind='Table'):
        follows_table = d._body._body[-2].tag == qn('w:tbl') if len(d._body._body) > 1 else False
        a = d.add_paragraph(t, style=kind + ' Caption')
        a.paragraph_format.first_line_indent = Inches(0)
        a.paragraph_format.keep_with_next = True
        if follows_table:
            a.paragraph_format.space_before = Pt(8)

    def table(headers, rows, widths=None):
        t = d.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = widths or [6.7 / len(headers)] * len(headers)
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
        for c, x in zip(t.rows[0].cells, headers):
            c.text = str(x).replace('all498', 'all 498')
        for row in rows:
            for c, x in zip(t.add_row().cells, row):
                text = re.sub(';(?=\\S)', '; ', str(x)).replace('score0,1or2 divided by2', 'score 0, 1 or 2 divided by 2')
                if isinstance(x, str) and len(text) > 10 and (not ('_' in text and ' ' not in text)) and (not re.fullmatch('[a-f0-9]{12,64}', text)):
                    text = re.sub('([A-Za-z])(?=\\d)', '\\1 ', text)
                    text = re.sub('(\\d)(?=[A-Za-z])', '\\1 ', text)
                    text = text.replace('S 0', 'S0').replace('S 1', 'S1')
                c.text = text
        for ri, row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            if ri == 0:
                pr.append(OxmlElement('w:tblHeader'))
            for ci, c in enumerate(row.cells):
                c.width = Inches(widths[ci])
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = c._tc.get_or_add_tcPr()
                bd = OxmlElement('w:tcBorders')
                for side in ['top', 'bottom', 'left', 'right']:
                    el = OxmlElement('w:' + side)
                    el.set(qn('w:val'), 'single')
                    el.set(qn('w:sz'), '4')
                    el.set(qn('w:color'), 'D9D9D9')
                    bd.append(el)
                pr.append(bd)
                m = OxmlElement('w:tcMar')
                for side, value in [('top', 90), ('bottom', 90), ('left', 100), ('right', 100)]:
                    el = OxmlElement('w:' + side)
                    el.set(qn('w:w'), str(value))
                    el.set(qn('w:type'), 'dxa')
                    m.append(el)
                pr.append(m)
                if ri == 0:
                    sh = OxmlElement('w:shd')
                    sh.set(qn('w:fill'), 'EDEFF0')
                    pr.append(sh)
                for a in c.paragraphs:
                    a.paragraph_format.first_line_indent = Inches(0)
                    a.paragraph_format.space_after = Pt(0)
                    a.paragraph_format.line_spacing = 1.05
                    a.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in a.runs:
                        r.font.size = Pt(10)
                        r.bold = ri == 0

    def pic(file, width=6.5):
        a = d.add_paragraph()
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a.paragraph_format.first_line_indent = Inches(0)
        a.add_run().add_picture(str(OUT / file), width=Inches(width))
    a = d.add_paragraph()
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.add_run().add_picture(str(LOGO), width=Inches(1.25))
    p('รายงานโครงงาน', True, True, 15)
    p(TH_TITLE, True, True, 12)
    a = p('(' + TITLE + ')', True, size=12)
    a.style = d.styles['Title']
    p()
    p('จัดทำโดย', True, True, 12)
    for name, sid, email in MEMBERS:
        p(name + '    รหัสนักศึกษา ' + sid, True, True, 11)
    p()
    p('อาจารย์ที่ปรึกษา', True, True, 12)
    p(ADVISOR, True, True, 11)
    p()
    p('รายงานนี้เป็นส่วนหนึ่งของการศึกษารายวิชาโครงงาน', True, True, 11)
    p(FACULTY + '  สาขา' + PROGRAMME, True, True, 11)
    p(INSTITUTION, True, True, 11)
    p(TERM, True, True, 11)
    page('Abstract')
    p('This project implements a document assistant for Thai business reports using OCR, table extraction, locally stored embeddings and structured numerical access. We compared a one-pass dense retrieval system (S0) with the selected full application (S1) on the same newly ingested reports, questions, final generation model, context budget and validators. Four 2025 issuer reports contained 1,447 physical PDF pages; 1,431 pages were indexed, with 16 failed checkpoints and additional table-quality warnings retained. Each system produced first outcomes for 498 answerable questions and 32 unsupported controls.')
    p('Strict-complete results are 79/498 (15.86%) for S0 and 60/498 (12.05%) for S1. Required numerical matches are 119/356 (33.43%) and 106/356 (29.78%). Appropriate control abstentions remain 24/32 and 28/32. Median response times remain 7.19 and 13.01 seconds. S1 did not outperform S0 on these strict criteria; it abstained more appropriately on controls and took longer.')
    p('All 1,060 retained first outcomes now have usable AI-assisted assessments. A separate completion overlay resolves the 348 earlier failed judgments using saved historical responses where available and current Codex review for missing decisions. No new Gemini calls or application answers were generated. Thirty-four assessment records received disclosed consistency adjudications. Original files and the earlier result manifest remain intact; 2,124 frozen evidence hashes were verified. Reviewer assignment was not randomized and no independent human confirmation is available. Three confirmed source-label defects are disclosed with a separate paired sensitivity analysis. The greater-than-80-percent quality goal remains unmet. Costs include known partial components and explicitly unknown charges.')
    p('Keywords: Thai business reports; OCR; retrieval augmented generation; numerical context; source provenance; AI-assisted evaluation.', size=10.5)
    d.add_heading('Acknowledgements', level=2)
    p('The project advisor and institutional information follow the supplied examination report. We acknowledge the original issuer reports as the evidence sources. AI assistance in document extraction, answer generation, evaluation and preparation of this report is disclosed in Chapter 3; it does not constitute independent human verification.')
    d.add_page_break()
    d.add_paragraph('Contents', style='Front Heading')
    a = d.add_paragraph()
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
    a._p.append(f)
    d.add_page_break()
    d.add_paragraph('List of Tables', style='Front Heading')
    a = d.add_paragraph()
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), 'TOC \\t "Table Caption,1" \\h \\z')
    a._p.append(f)
    d.add_paragraph('List of Figures', style='Front Heading')
    a = d.add_paragraph()
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), 'TOC \\t "Figure Caption,1" \\h \\z')
    a._p.append(f)
    page('Chapter 1 Introduction')
    h('1.1 Problem and intended workflow', 'Business-report research requires more than locating a number. An analyst must identify the entity, year, unit, consolidation scope and meaning of the figure, then return to the original page. Thai reports combine narrative sections with dense tables and financial notes. A plausible answer can still be unsuitable when it uses a liability instead of an asset, a target instead of an actual result, or a figure from another year. The prototype is intended for a single operator or small research team that repeatedly consults company reports and retains source evidence for later inspection.')
    h('1.2 Objectives', 'The project aims to collect inspectable document sources, extract searchable prose and structured table cells, retrieve numerical and narrative evidence, and generate answers with captured claims, numerical tuples and physical PDF citations. The evaluation objective is to compare the selected full application with a simpler dense baseline under common inputs. A quality target above80percent was an objective, not a result; the final measured run did not meet it.')
    h('1.3 Scope and boundaries', 'The official project title retains Web Scraping. In this final measured study, the evidence corpus consists of four downloaded official PDF reports, not a benchmark of arbitrary website extraction. Storage, SQL access and bge-m3 embeddings run locally; historical OCR and generation use online Typhoon and Gemini services. The workflow is therefore not fully offline. Knowledge-graph and Tavily functionality were not required in the selected closure build. No claim of enterprise multi-user readiness, commercial superiority or autonomous financial decision-making is established.')
    h('1.4 Research questions', 'We ask whether the complete application improves strict completion and numerical-relation matches over one-pass dense retrieval; whether it handles unsupported controls appropriately; and what additional response time, tool activity and known-token components it incurs. Source spot checks separately ask whether facts, captured annotations and actual citations agree with the original PDF. These checks are diagnostic and do not retroactively tune the frozen experiment.')
    page('Chapter 2 Background and Related Work')
    h('2.1 Retrieval augmented generation', 'Retrieval augmented generation combines a language generator with retrieved external evidence. Lewis et al. [1] studied dense retrieval with external memory for knowledge-intensive tasks. The prototype follows this broad separation: searchable report content supplies document-specific evidence to a final generator. Retrieval does not itself guarantee correctness; relevant passages can be omitted and extracted table headings can be incomplete.')
    h('2.2 Reasoning and tool use', 'ReAct [2] combines reasoning and actions that obtain information from external tools. Here, the full application routes requests to SQL or text retrieval and may attempt search repair. The final study compares entire systems, not an isolated lexical-ranking or SQL ablation. Tool registration is not evidence of use; Chapter 4 reports captured actions. The shared generator and common validators reduce some differences but do not make the systems identical in retrieval behaviour.')
    h('2.3 Official product survey', 'Official product documentation shows that commercial assistants can accept documents and project sources. ChatGPT Plus provides file upload and analysis [6]; Claude Pro includes document-oriented chats and projects [7]; Google AI Pro bundles Gemini services [8]; Gemini Notebook supports PDF and other source formats [9]. This desk survey establishes available features only. No commercial assistant was run on this question bank, so relative accuracy, usability and savings remain untested. The project contribution is its specific Thai numerical context and inspectable experiment records, rather than exclusive access to specialized documents.')
    h('2.4 Numerical provenance', 'A numerical relation includes entity, measure, value, unit/scale, year, qualifier and source page. A subtraction also needs correctly bound operands and a direction. Physical PDF numbering is used throughout this report; it can differ from printed folios. Original issuer reports [10-13] are used to audit extraction and citations. Human adjudication remains necessary where rounding precision, date duration or reference breadth is ambiguous.')
    page('Chapter 3 System and Experimental Method')
    h('3.1 Implemented processing architecture', 'Official PDFs are retained with source URLs and SHA-256 identifiers. The main ingestion route uses Typhoon for page prose and Gemini for table extraction. PostgreSQL with pgvector stores searchable content and page checkpoints; DuckDB stores extracted table cells. Local bge-m3 embeddings have1,024 dimensions. The final eligible corpus contains6,928 chunks. These records remain machine derived: an indexed page can still contain table or semantic extraction errors.')
    cap('Figure 3.1 Shared processing workflow and separated evaluation evidence', 'Figure')
    pic('workflow.png')
    h('3.2 Matched systems', 'S0 performs one-pass dense retrieval of ten distinct pages with common page expansion. S1 uses the selected application with routing, SQL, lexical/vector retrieval, search repair and deterministic cell handling. Graph-catalog functionality is disabled. Both systems use the same actual OCR corpus, Gemini2.5Flash final generator, final-answer prompt,32,000-character context budget and validators. Source labels and native PDF text are excluded from generation and used only for source evaluation. This is a whole-system comparison; results cannot identify which individual tool caused a difference.')
    h('3.3 Inspection workflow', 'Answers retain text, captured claims, numeric facts, source indices, claim-to-source links and tool traces. Source inspection begins from the emitted document and physical page, checks the original table or paragraph, and confirms the requested entity and numerical dimensions. A truthful citation to a different measure does not make an answer relevant. The live status viewer reads checkpoints only and makes no model requests.')
    page('3.4 Reports questions and freeze', level=2)
    cap('Table 3.1 Locked source reports and final answerable questions')
    table(['Issuer', 'Physical pages', 'Indexed pages', 'Questions'], [['KCE', 329, 329, 125], ['OSP', 375, 371, 124], ['HMPRO', 470, 462, 125], ['TISCO Bank', 273, 269, 124], ['Total', 1447, 1431, 498]], [2.2, 1.5, 1.5, 1.5])
    p('The final answerable bank contains199 numerical lookups,100 comparison/multiple-fact questions and199 narrative/policy/ESG/governance questions. Earlier planned category totals differ from these final locked counts. Each report also has eight unsupported controls, yielding32 controls per system. The total is530 outcomes per system and1,060 across both. Shared question IDs permit paired analysis; two systems answering the same question are not independent question samples.')
    p('Four issuer reports provide broad page coverage but remain a small document sample. They were selected as new final reports rather than reused development outcomes. Source references and numerical requirements were prepared with AI assistance and locked before final generation. This is held-out application measurement under that process, not a human-established benchmark. Earlier development scores are excluded from the final tables and are not relabelled as unseen results.')
    h('3.5 Frozen first outcomes', 'The source bank, labels, application code, configuration, prompt, scorer and analysis were frozen for the run. First outcomes, capture errors and refusals were retained. No completed wrong answer was retried to select a passing answer. Operational scheduling and recovery wrappers were disclosed separately. Original source documents and frozen experiment files remain preserved; this report and supplementary QA introduce no database or scorer edits.')
    h('3.6 Data separation', 'Native source text and page images may resolve evaluation ambiguity but must not replace the actual OCR generation corpus. The frozen source-label SHA-256 begins 5959f4e2e4cc. File hashes and record inventories in the appendices support traceability. Hash agreement proves the referenced bytes remained the same; it does not prove that a reference claim or AI judgment is correct.')
    page('3.7 Evaluation definitions', level=2)
    p('Strict-complete success is deliberately demanding. A non-abstaining, error-free captured answer must cover all minimal reference claims, have supported/relevant factual claims, match all required numerical relations, leave no unannotated assertions, and provide at least one fully supporting actual citation per claim. Its prose must explicitly state physical PDF pages consistent with the supported citation pages. A correct narrative with structured citations but no page wording can therefore fail. This metric should not be described as ordinary accuracy.')
    cap('Table 3.2 Custom question-level metric definitions')
    table(['Metric', 'Definition'], [['Faithfulness', 'Fraction of captured answer claims supported by actual generation context'], ['Factual precision', 'Fraction of answer claims supported by the original source'], ['Factual recall', 'Fraction of minimal required source claims covered by the answer'], ['Factual F1', 'Harmonic mean of question-level factual precision and recall'], ['Claim relevance precision', 'Fraction of captured claims relevant to the question'], ['Answer relevancy rubric', 'Judgment score0,1or2 divided by2'], ['Context precision AP', 'Average precision of useful ordered context blocks'], ['Context recall', 'Fraction of required claims covered by actual context']], [2.1, 4.6])
    p('These are custom source-judgment metrics, not scores produced by the Ragas library. The report shows assessable means with their denominators and conservative means over all498 main questions. Missing or non-applicable question values contribute zero to conservative means. They are not silently dropped from whole-set results. Claim precision is a mean of question-level fractions, not a claim-weighted micro-average. Numerical rates use356 required relations per system; controls use their separate32-question denominator.')
    h('3.8 Paired uncertainty', 'S1-minus-S0 differences use matched main questions. The frozen analysis resamples four issuer-report clusters with2,000 bootstrap replicates and seed20261009. This preserves report grouping better than treating every question as independent, but four clusters provide limited generalization. Missingness, reviewer differences and detected judgment inconsistencies limit interpretation of confidence intervals. The intervals do not establish causal superiority or correct an unreliable judgment.')
    page('3.9 Execution recovery and AI reviewers', level=2)
    p('Two unexpected host shutdowns occurred during processing: one during OCR and one after643 saved generation outcomes. Durable checkpoints and outcome inventories allowed continuation without replaying completed outcomes. The original runner had a resume-only Python scope error. Its source file stayed unchanged; a disclosed wrapper inserted only a global freeze declaration, and removing the declaration restored the original AST. This repaired runtime name lookup without changing questions, corpus, generation or scorer logic.')
    p('The failed original resume had already opened DuckDB through existing schema initialization, changing its physical storage hash. A comparison against an untouched warehouse copy including its original WAL verified identical table columns, row counts, canonical row-content hashes and view definitions after the same initialization. The earlier comparison without WAL was incomplete and excluded. Frozen vectors and PostgreSQL corpus identifiers were also checked. Physical byte identity and logical equivalence are distinct claims.')
    p('After generation, source-judgment scheduling increased from a sequential worker to disjoint parallel shards, eventually five at the user’s request. Prompts, parameters and validation stayed frozen; existing judgments and SDK attempts remained retained. A coordinator later stopped with a OneDrive permission error while owned workers continued. The user then explicitly stopped Gemini. All owned judge processes were stopped, including two in-flight requests whose final provider status and usage remain unknown. No further Gemini calls were made for Codex review or artifact preparation.')
    cap('Table 3.3 Reviewer provenance and assessment availability')
    table(['Records', 'Count', 'Interpretation'], [['Previously usable assessments', 712, '604 historical Gemini; 108 Codex'], ['New completion overlays', 348, 'Saved components plus Codex decisions'], ['Consistency adjudications', 34, 'Separate layer; originals preserved'], ['All usable assessments', 1060, '0 unassessed; 0 human-confirmed']], [2.4, 1.0, 3.3])
    p('The completed edition retains 712 previously usable assessments and adds 348 separate completion overlays. Existing saved provider responses supplied recoverable components; Codex evaluated missing decisions against the original evidence and frozen validator. A subsequent audit applies 34 versioned consistency adjudications, including prior positive numerical decisions that violated exact quantity, sign, scope or captured-annotation requirements. No original answer or assessment is overwritten. Reviewer assignment was not randomized or blind; source documents and question identities were visible. All judgments remain AI-assisted, including visually inspected source pages. Schema validity does not establish semantic correctness.')
    page('Chapter 4 Results and Evidence')
    p('All 1,060 first outcomes are retained and all 1,060 assessment records are usable. There are zero unassessed judgments. Tables below come from RESULT_MANIFEST_v2 using the unchanged scorer and frozen primary references. A canonical input projection changes 16 null annotation collection fields to empty lists without inventing claims, tuples or context. All original outcome files remain byte-identical. Completed assessment coverage is distinct from answer success and from metric applicability.')
    cap('Table 4.1 Completed primary results and response time')
    table(['Measure', 'S0 Dense RAG', 'S1 Full App'], [['Strict complete /498', f'79 /498 ({pct(S0['complete_answer_rate_all498'])})', f'60 /498 ({pct(S1['complete_answer_rate_all498'])})'], ['Numerical relations /356', f'119 /356 ({pct(S0['numeric_rate'])})', f'106 /356 ({pct(S1['numeric_rate'])})'], ['False refusals /498', 158, 176], ['Capture-error questions /498', 34, 32], ['Model-error outcomes', 0, 0], ['Unassessed main judgments', 0, 0], ['Appropriate control abstention /32', '24 /32 (75.00%)', '28 /32 (87.50%)'], ['Truthful irrelevant controls /32', 8, 4], ['Unsupported requested control assertions', 0, 0], ['Observed wall time p50', '7.19 s', '13.01 s'], ['Observed wall time p95', '19.90 s', '35.55 s']], [3.1, 1.8, 1.8])
    p('Strict completion differs by -3.82 percentage points for S1 relative to S0. The four-report bootstrap interval is -6.64 to -1.20 points. The numerical rate differs by -3.65 points. More SQL activity did not establish more successful numerical relations. These are paired observations from four reports, with AI-assisted judgments and limited generalization; they do not establish a causal explanation.')
    p('Controls ask for unsupported constructed entities or future actuals. The remaining control answers substituted truthful report-supported historical figures that did not answer the request. They are counted as irrelevant substitutions, separately from unsupported assertions about the requested entity/year. Better abstention on controls does not offset false refusals on answerable questions.')
    page('4.2 Answer and retrieval metrics', level=2)
    cap('Table 4.2 Assessed means and conservative means for eight custom metrics')
    metricnames = {'faithfulness': 'Faithfulness', 'factual_precision': 'Factual precision', 'factual_recall': 'Factual recall', 'factual_f1': 'Factual F1', 'claim_relevance_precision': 'Claim relevance precision', 'answer_relevancy_rubric': 'Answer relevancy rubric', 'context_precision_ap': 'Context precision AP', 'context_recall': 'Context recall'}
    table(['Metric', 'S0 assessed\nmean / N', 'S1 assessed\nmean / N', 'S0 all498', 'S1 all498'], [[metricnames[k], f'{pct(S0['rag_metrics'][k]['assessed_mean'])} / {S0['rag_metrics'][k]['assessable_n']}', f'{pct(S1['rag_metrics'][k]['assessed_mean'])} / {S1['rag_metrics'][k]['assessable_n']}', pct(S0['rag_metrics'][k]['conservative_all498_mean']), pct(S1['rag_metrics'][k]['conservative_all498_mean'])] for k in metricnames], [2.05, 1.45, 1.45, 0.875, 0.875])
    p('Assessed-only factual precision is 96.38% for S0 and 92.85% for S1, based on 336 and 318 applicable question values. These are not whole-bank accuracy rates. Refusals and uncaptured claims can have no claim-precision value even though their assessment is complete. The conservative columns retain all 498 questions and assign zero to missing or non-applicable metric values.')
    p('Conservative context precision AP is 47.10% versus 52.32%; context recall is 38.05% versus 41.12%. Paired report-cluster intervals are retained in the result manifest. Retrieval averages do not guarantee complete answers when extraction, binding, capture or refusal guards fail.')
    cap('Figure 4.1 Recorded primary rates with separate denominators', 'Figure')
    pic('primary_results.png', 5.8)
    page('4.3 Results by category and report', level=2)

    def groups(key):
        result = []
        names = {'numerical_lookup': 'Numerical lookup', 'comparison_multifact': 'Comparison / multiple facts', 'narrative_policy_esg_governance': 'Narrative / policy / ESG / governance', 'TISCO_BANK': 'TISCO Bank'}
        for value in sorted(set((r[key] for r in ROWS if r['answerable']))):
            a = [r for r in ROWS if r['answerable'] and r[key] == value and (r['system'] == 'simple_dense_rag')]
            b = [r for r in ROWS if r['answerable'] and r[key] == value and (r['system'] == 'full_app')]
            result.append([names.get(value, value), len(a), f'{sum((r['complete_answer'] for r in a))}/{len(a)}', f'{sum((r['complete_answer'] for r in b))}/{len(b)}', sum((r['judge_status'] != 'first_validated_judgment_retained' for r in a)), sum((r['judge_status'] != 'first_validated_judgment_retained' for r in b))])
        return result
    cap('Table 4.3 Main question categories and retained strict-complete counts')
    table(['Category', 'N each', 'S0 complete', 'S1 complete', 'S0 missing', 'S1 missing'], groups('category'), [2.3, 0.7, 0.95, 0.95, 0.9, 0.9])
    cap('Table 4.4 Issuer report results and judgment missingness')
    table(['Report', 'N each', 'S0 complete', 'S1 complete', 'S0 missing', 'S1 missing'], groups('issuer'), [2.3, 0.7, 0.95, 0.95, 0.9, 0.9])
    p('There are no missing assessment records in this edition. Genuine non-applicable metric values remain visible. Historical reviewer assignment was concentrated by report and category, and the new completion overlay contains both recovered historical components and Codex decisions. Apparent stratum differences can combine system behaviour, source complexity, extraction quality and assessor effects. Both systems retain the same original outcomes and question denominators.')
    h('4.4 Source QA design', 'One original first outcome was selected from each system/report/category stratum, irrespective of success, using seed20261009. The resulting24 spot checks inspect selected original source facts, actual captured numerical tuples and emitted citations. They are supplementary observations, not exhaustive regrading of every claim. Page images and notes are retained with original outcome/judgment hashes. Every stratum is represented; the sample does not certify the entire bank.')
    page('4.5 Numerical and citation audit findings', level=2)
    p('The completion audit found inconsistent positive numerical decisions in historical judgments. For example, KCE03q05 S0 reports source-supported 263,887,772 baht against a frozen 263.9-million reference. This finer-precision quantity is not an exact equality to that reference after unit scaling, so the strict numerical flag is corrected to false in a separate adjudication. Source-supported finer precision is not thereby a hallucination. Semantic reference coverage and factual support are separate criteria. No favorable rounding tolerance was introduced after seeing results.')
    p('TISCO06q15 S1 originally marked a five-year answer as a numerical match while its reason called it a mismatch against a four-year subtraction. The source gives BE2564 to December BE2568 without a start month; exact elapsed duration is unresolved. The versioned adjudication marks exact-duration support insufficient and the required numerical relation a nonmatch. The original contradictory assessment is preserved. Uniform date-only treatment is also documented for the corresponding paired record and the partial-month tenure question.')
    p('Numerical capture and factual text can also diverge. TISCO05q10 S0 correctly states the Stage1 twelve-month ECL policy at cited page202 but emits no numeric facts. TISCO06q07 S0 correctly states the tenure dates at cited page174 but emits no start/end tuples. The frozen numerical criterion fails these captured annotations. For OSP04q08 S0, the answer states an increase of one event while tuples contain only operands2and1; the retained judge credits the derivation through operand indices. Annotation and derivation treatment need calibration before a new benchmark.')
    p('Measure identity is a separate failure. TISCO03q03 S1 returns142,967thousand baht of long-term employee-benefit expense from page239, whereas the request’s broader employee compensation total is2,569,856,472.52baht on page92. HMPRO05q15 S0 gives truthful lease expiry/renewal details from page72 but omits the requested recognition policy on page416. Truthful numbers or paragraphs do not by themselves answer the question.')
    h('4.6 Frozen source labels and paired sensitivity', 'Three source-label defects were confirmed against original reports. KCE01q09 compares Thai laminate/prepreg sales with a Taiwan amount that also includes copper foil. OSP08q12 uses two same-year reporting-scope columns as though they were prior/current years. KCE08q02 asks total separate long-term borrowing but its zero reference is only the fixed-rate subgroup; the floating-rate and balance-sheet tables support 238,100 thousand baht. Primary references remain frozen. These questions are excluded symmetrically only in the following diagnostic sensitivity.')
    cap('Table 4.5 Paired source-validity sensitivity excluding three defective questions')
    sx=M['source_validity_sensitivity']['systems']
    table(['Criterion', 'S0 Dense RAG', 'S1 Full App'], [['Strict complete /495',str(sx['simple_dense_rag']['strict_complete'])+'/495 ('+pct(sx['simple_dense_rag']['strict_rate'])+')',str(sx['full_app']['strict_complete'])+'/495 ('+pct(sx['full_app']['strict_rate'])+')'], ['Numerical relations /353',str(sx['simple_dense_rag']['numeric_correct'])+'/353 ('+pct(sx['simple_dense_rag']['numeric_rate'])+')',str(sx['full_app']['numeric_correct'])+'/353 ('+pct(sx['full_app']['numeric_rate'])+')']], [3.1,1.8,1.8])
    p('These sensitivity results are not a repaired or newly unseen benchmark. They expose known source-label defects without changing primary denominators. Broader ambiguities include reporting basis, requested versus extra reference facts, partial-date duration and exact rounded precision. Correct prose can still fail if the required captured numeric tuple is absent. Explicit signed decreases may be equivalent to a negative change; a generic payment row is not silently sign-flipped. All adjudications and their reasons remain separately identifiable.')
    h('4.7 Traceable cash example', 'TISCO Bank page128 displays BE2568 cash as818,057thousand baht; page266 displays rounded818million. The document, entity, year, scale and requested precision must remain attached to the value. The source check confirms these displays, not a new successful application outcome. This example is used consistently in the poster. Committee duties across page breaks similarly require the preceding heading to avoid assigning an executive duty to the risk committee.')
    page('4.8 Ingestion reliability and actual tool activity', level=2)
    p('Final indexing covered1,431/1,447 physical pages (98.89percent). This is checkpoint coverage, not OCR accuracy. Failed physical pages were OSP8,21,277,320; HMPRO1,174,191,295,297,321,461,468; and TISCO79,206,228,229. KCE had all329 pages indexed but retained table-quality warnings. Unresolved rows and partial semantic extraction problems remain in the corpus and were not relabelled as successful human verification.')
    cap('Table 4.6 Captured tool actions across 530 outcomes per system')
    tools = sorted(set(S0['actual_tools']) | set(S1['actual_tools']))
    table(['Captured action', 'S0 count', 'S1 count'], [[k, S0['actual_tools'].get(k, 0), S1['actual_tools'].get(k, 0)] for k in tools], [3.5, 1.6, 1.6])
    p('Tool counts are action occurrences, not necessarily distinct questions. S1 used SQL314 times and vector search435 times, with seven self-correction and22 multi-hop actions. The multi-hop action label does not establish knowledge-graph use. S0 used vector search530 times. More tool attempts also created more provider calls, which are separated from final outcome counts in cost accounting.')
    p('Observed latency includes the time associated with each retained outcome. It is distinct from total elapsed project duration and shutdown downtime. Median S1 wall time was about1.81times S0; p95 was about1.79times. These observations describe this machine, provider state and workload. They are not guaranteed service-level performance or a universal parallel throughput limit.')
    h('4.9 Retained failures', 'Capture errors occurred in34 S0 and32 S1 main questions. No final outcome was classified as a model error by the frozen summary, but this does not imply zero provider attempt errors. The final generation journals contain2,322 SDK attempts:2,079 successes and243 errors, including234 recorded429 errors. Provider retries already present in the frozen runner remain in the accounting; completed outcomes were not rerun to improve scores.')
    page('4.10 Operating cost and unknown costs', level=2)
    p('Costs are filtered to the four locked PDF hashes for ingestion and the exact final paired run ID for answer generation. Historical calibration calls are excluded from operation estimates. The known Gemini component uses standard uncached input at USD0.30per million tokens and output including reported thinking at USD2.50per million, checked against official pricing [3] on11October2026. Cache input is cheaper, but actual cache use is unknown. These components are not total invoices or guaranteed lower bounds: cache could reduce input cost while missing usage/thinking could increase totals.')
    cap('Table 4.7 Filtered operation attempts and known Gemini components')
    table(['Operation', 'Provider attempts', 'Known Gemini USD', 'Key unknowns'], [['Four-report ingestion', '2,542 total\n1,065 Gemini;1,477 Typhoon', '3.436368', '64 Gemini usage gaps; thinking unreported'], ['S0 queries /530 outcomes', 666, '4.074021', '54 usage / thinking gaps'], ['S1 queries /530 outcomes', 1656, '5.610250', '189 usage gaps;934 thinking gaps']], [2.1, 1.55, 1.3, 1.75])
    p('Typhoon’s official hosted model page advertises a free access tier [4], but the actual account tier and invoice are not established here. Codex review account allocation, hardware depreciation, electricity, administration, tax and credits are also unknown. Interrupted requests with no usage remain unknown. The retained closure research has a separate known Gemini component of USD29.541234 at cost checkpointv7; this includes source work and earlier research, and is excluded from monthly operation scenarios.')
    p('The accounting retains every SDK attempt rather than counting only successful final answers. Ingestion attempts include reprocessing and provider failures already recorded during the original run. Parallel source review changes waiting time but does not make the evaluation cost disappear. The historical research subtotal is separate from answering on a prepared corpus. Account-level charges cannot be allocated to this project without billing evidence; no numerical value is assigned to unreported Codex usage or local energy.')
    page('4.11 Illustrative monthly workload', level=2)
    p('For a prepared corpus, the recorded known component per outcome is USD0.00768683 for S0 and USD0.01058538 for S1. A report of similar average length contributes an illustrative ingestion component of USD0.859092. Using Q monthly questions and N new reports, the arithmetic is Q times the per-outcome component plus N times0.859092. This assumes similar report length and behaviour, not measured production scaling.')
    cap('Table 4.8 Illustrative monthly known-token components in USD')
    scenarios = [(50, 0), (200, 0), (1000, 0), (200, 1), (1000, 4)]
    table(['Questions Q', 'New reports N', 'S0 component', 'S1 component'], [[q, n, f'{q * COST['systems']['simple_dense_rag']['known_component_per_first_outcome_USD'] + n * 0.859092:.2f}', f'{q * COST['systems']['full_app']['known_component_per_first_outcome_USD'] + n * 0.859092:.2f}'] for q, n in scenarios], [1.5, 1.5, 1.85, 1.85])
    p('Subscription products offer broader services and have different limits and account terms. Partial API components cannot establish subscription-equivalent savings, a break-even threshold or return on investment. A fair economic comparison would need real workload, account invoices, maintenance time and a measured quality/utility comparison.')
    p('The source reports average361.75physical pages. Reports with more tables, failed OCR pages or longer retrieved contexts may consume more attempts and tokens than this arithmetic assumes. Cache use and provider pricing can also change. The monthly examples are planning scenarios based on the recorded component, not promises of a production bill. Repeated queries were not rerun here to measure cache savings or concurrent response capacity.')
    cap('Table 4.9 Official subscription price examples for desk comparison')
    table(['Product', 'Listed monthly USD', 'Comparison scope'], [['ChatGPT Plus [6]', '20.00', 'File analysis; API billed separately'], ['Claude Pro [7]', '20.00', 'Monthly subscription; annual option differs'], ['Google AI Pro [8]', '19.99', 'US monthly listing; broader Google bundle']], [2.1, 1.55, 3.05])
    p('These official listings were checked on11October2026. They do not establish Thai checkout prices, taxes, product quotas for this workload or comparative answer quality. No subscription was purchased or tested as part of this final study. The prototype’s inspectable storage and source records are potential workflow benefits, but a useful comparison must include the time spent checking unreliable answers.')
    page('Chapter 5 Conclusions and Future Work')
    h('5.1 Findings', 'The project produced an inspectable Thai business-report processing and answer workflow and completed the paired first-outcome run. The selected full application did not outperform one-pass dense RAG on the preserved strict-complete and numerical criteria. It handled unsupported controls more appropriately but increased median and tail response times. Slightly higher retrieval averages did not resolve answer binding, refusal or annotation problems. The greater-than-80-percent quality goal remains unmet.')
    h('5.2 Objectives and practical use', 'Document collection, durable page processing, local evidence storage, numerical access and source inspection were demonstrated. Automatic answer reliability was not demonstrated at the target level. The prototype can support supervised evidence discovery, with the operator opening source pages and checking entity, period and units. It should not present generated figures as verified business conclusions merely because a document citation is attached.')
    h('5.3 Limitations', 'Only four reports were evaluated. Sixteen failed page checkpoints and table warnings restrict the corpus. All assessments are complete, but source labels and judgments remain AI-assisted and nonrandomly assigned. Three confirmed source-label defects and other scope, precision and duration ambiguities are disclosed. Strict page-wording and captured-tuple requirements can reject factually supported text. Unknown costs prevent a complete economic comparison.')
    h('5.4 Next research steps', 'A new evaluation should begin with independent human adjudication of reference scope, rounding tolerances and duration definitions, followed by calibrated semantic judgment checks. Retrieval and capture issues should be corrected in a new explicitly versioned build, then compared with the baseline on a newly frozen held-out set. Failed OCR pages and ambiguous table headers need systematic repair and verification. Any later commercial comparison should run the same permissible document tasks under documented product settings and report uncertainty. These proposals are future work, not completed improvements or permission to retune this frozen run.')
    h('5.5 Reporting commitment', 'All primary values in this report and poster come from the completed versioned result manifest. Original first outcomes, original judgments and the v1 manifest remain preserved. Completion overlays and 34 assessment adjudications are distinct from system improvements. Paired source-validity sensitivity does not replace frozen primary results. Original project, member, advisor, institution and academic-term identities remain unchanged.')
    page('References')
    refs = ['[1] Lewis, P., et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. https://arxiv.org/abs/2005.11401', '[2] Yao, S., et al. (2023). ReAct Synergizing Reasoning and Acting in Language Models. https://arxiv.org/abs/2210.03629', '[3] Google Cloud. Gemini model pricing. Official standard Gemini2.5Flash input/output rates, verified11October2026. https://cloud.google.com/gemini-enterprise-agent-platform/generative-ai/pricing', '[4] Typhoon. Typhoon OCR hosted model and access information. https://opentyphoon.ai/model/typhoon-ocr', '[5] Google Cloud. Throughput and quota for generative AI. https://docs.cloud.google.com/vertex-ai/generative-ai/docs/resources/throughput-quota', '[6] OpenAI. What is ChatGPT Plus. https://help.openai.com/en/articles/6950777-what-is-chatgpt-plus', '[7] Anthropic. Claude pricing and Pro features. https://claude.com/pricing', '[8] Google. Google AI plans. https://one.google.com/about/google-ai-plans/?hl=en&gl=US', '[9] Google. Gemini Notebook supported source formats. https://support.google.com/gemininotebook/answer/16215270?hl=en']
    for i, r in enumerate(BANK['reports'], 10):
        refs.append(f'[{i}] {r['issuer'].replace('_', ' ')}. One Report2025 Thai edition. Original retained PDF, {r['pages']} physical pages. {r['source_url']}')
    refs.append('[14] Project frozen first-outcome records, source labels, RESULT_MANIFEST_v2, filtered cost components and supplementary QA. Local evidence index in Appendix B. AI-assisted records; not independent human validation.')
    for ref in refs:
        ref = ref.replace('Gemini2.5Flash', 'Gemini 2.5 Flash').replace('verified11October2026', 'verified 11 October 2026').replace('One Report2025', 'One Report 2025')
        a = p(ref, size=10.5)
        a.alignment = WD_ALIGN_PARAGRAPH.LEFT
    page('Author Profiles')
    for i, (name, sid, email) in enumerate(MEMBERS, 1):
        d.add_heading('สมาชิกคนที่ ' + str(i), level=2)
        for text in ['ชื่อ-สกุล: ' + name, 'รหัสนักศึกษา: ' + sid, 'สาขาวิชา: ' + PROGRAMME, 'คณะ: เทคโนโลยีสารสนเทศ', 'สถาบัน: ' + INSTITUTION, 'อีเมล: ' + email]:
            p(text, thai=True, size=11)
        p()
    page('Appendix A Supplementary Source Spot Checks')
    p('The original 24 stratified samples cover both systems, four reports and three categories. Their dated findings were diagnostic observations from the earlier assessment checkpoint. This edition separately completes all missing judgments and applies disclosed consistency corrections; original sample notes remain preserved. The sample does not certify all claims or confer human confirmation.')
    cap('Table A.1 Stratified source samples and observations')
    short = ['Bond difference -1116; refusal', 'Paris claim supported; null units', 'Contractor LTIFR0.32; refusal', 'Revenue decrease1871.8; refusal', 'Appointment source164; refusal', '15 suppliers and citation agree', 'ROE change12.39pp; refusal', 'Selected teamwork claim supported', 'Lease carrying158366; source supported', 'Duration ambiguity; adjudicated', 'Revaluation policy/citation agree', 'Employee measure mismatch', 'Income difference -2006.74; refusal', 'Truthful lease details; wrong policy', 'Rounded cash rejection inconsistency', 'Singapore cash -6486; refusal', 'Wrong-company retrieval; refusal', 'Fine precision; strict nonmatch', 'Correct operands; derived capture issue', 'Selected human-rights claims supported', '17 directors; valid alternate page', '7 attendances vs8 regular meetings', '12-month policy; missing tuple', 'Correct tenure text; missing tuples']
    table(['Record ID', 'Arm', 'Original pages', 'Selected observation'], [[s['filename'].replace('__full_app.json', '').replace('__simple_dense_rag.json', ''), 'S1' if s['system'] == 'full_app' else 'S0', ','.join(map(str, s['physical_original_pages_inspected'])), short[i]] for i, s in enumerate(QA['samples'])], [2.15, 0.5, 0.9, 3.15])
    # Compact the 24-row appendix enough to keep its final row with the sample table.
    for row in d.tables[-1].rows:
        for cell in row.cells:
            margin=cell._tc.get_or_add_tcPr().find(qn('w:tcMar'))
            for side in ('top','bottom'):
                margin.find(qn('w:'+side)).set(qn('w:w'),'50')
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:run.font.size=Pt(9.5)
    page('Appendix B Reproducibility and Evidence Index')
    cap('Table B.1 Original PDF identities')
    table(['Report', 'SHA-256 prefix', 'Physical pages'], [[r['issuer'], r['source_sha256'][:20], r['pages']] for r in BANK['reports']], [2.0, 3.3, 1.4])
    p('All paths below are relative to TestFile/project_closure_2026-10-09. RESULT_MANIFEST_v2 is the sole completed primary summary; v1 remains a historical checkpoint. Completion provenance identifies each retained assessment, overlay and adjudication. The frozen scorer was replayed with four path constants substituted; numerical and metric algorithms were unchanged. Canonical projection only converts null annotation collections to empty lists.')
    for label, path in [('Primary results', 'RESULT_MANIFEST_v2.json'), ('Per-question rows', 'RESULTS_BY_QUESTION.json'), ('Filtered cost components', 'FILTERED_COST_COMPONENTS_v1.json'), ('Locked source labels', 'SOURCE_REPORTS.json'), ('First outcomes', 'final_paired_v24_v1/first_outcomes/'), ('Original and Codex judgments', 'final_source_judgments_v1/judgments/'), ('Completed reviewer provenance', 'evaluation_completion_v2/REPLAY_PROVENANCE.json'), ('Assessment overlays and corrections', 'evaluation_completion_v2/'), ('Integrity audit', 'evaluation_completion_v2/INTEGRITY_AND_COVERAGE_AUDIT.json'), ('24 source spot checks', 'SOURCE_SPOT_CHECK_INDEX.json'), ('Original952 hash inventory', 'codex_source_review_v1/ORIGINAL_952_JUDGMENTS_SHA256.json'), ('Stop and reviewer change', 'codex_source_review_v1/GEMINI_STOP_AND_REVIEWER_CHANGE.json'), ('Runtime binding repair', 'reboot_recovery_20261010_v3/RESUME_BINDING_FIX_DISCLOSURE.json'), ('Warehouse logical proof', 'reboot_recovery_20261010_v2/WAREHOUSE_WITH_WAL_STANDARD_INIT_EQUIVALENCE.json')]:
        a = p(label.replace('Original952', 'Original 952') + ': ' + path, size=9.5)
        a.paragraph_format.first_line_indent = Inches(0)
        a.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p('Primary manifest SHA-256: ' + sha(ROOT / 'RESULT_MANIFEST_v2.json'), size=9)
    p('The final Word source, PDF, poster SVG/HTML, print PDF and preview PNG are generated locally without new Gemini calls. Word contents and lists are updated during native rendering. Visual QA checks every rendered report page and the complete A1 poster. This artifact QA checks presentation; it does not change the false quality-goal flag or confer human-confirmed experimental validity.')
    d.save(OUT / 'examination_report.docx')
    (OUT / 'ARTIFACT_SOURCE_MANIFEST.json').write_text(json.dumps({'primary_manifest_sha256': sha(ROOT / 'RESULT_MANIFEST_v2.json'), 'source_labels_sha256': M['source_labels_sha256'], 'supplementary_qa_sha256': sha(ROOT / 'SOURCE_SPOT_CHECK_INDEX.json'), 'quality_goal_completed': False, 'human_confirmed': False, 'gemini_called_for_artifacts': False, 'identity_preserved': True}, indent=2) + '\n', encoding='utf-8')
    print('Final examination Word source saved; render and visual QA required before delivery.')
if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == 'assets':
        assets()
    else:
        report()
